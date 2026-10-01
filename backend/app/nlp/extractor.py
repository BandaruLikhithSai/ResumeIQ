"""
Document text extraction.
Supports PDF, DOCX, TXT.  Falls back to OCR if pdfminer returns minimal text.
"""
import io
import os
import re
import logging

logger = logging.getLogger(__name__)


def extract_text(file_path: str, file_type: str) -> dict:
    """
    Extract raw text from a document.

    Returns:
        {
            "text": str,
            "ocr_used": bool,
            "error": str | None,
            "page_count": int
        }
    """
    handlers = {
        "pdf": _extract_pdf,
        "docx": _extract_docx,
        "txt": _extract_txt,
        "png": _extract_image,
        "jpg": _extract_image,
        "jpeg": _extract_image,
    }

    handler = handlers.get(file_type.lower())
    if not handler:
        return {"text": "", "ocr_used": False, "error": f"Unsupported file type: {file_type}", "page_count": 0}

    try:
        return handler(file_path)
    except Exception as exc:
        logger.exception("Text extraction failed for %s", file_path)
        return {"text": "", "ocr_used": False, "error": str(exc), "page_count": 0}


# ── PDF ─────────────────────────────────────────────────────────────────────

def _extract_pdf(file_path: str) -> dict:
    text = _pdf_via_pdfminer(file_path)

    # If extracted text is suspiciously short, try OCR
    if len(text.strip()) < 100:
        ocr_text = _try_ocr_pdf(file_path)
        if len(ocr_text.strip()) > len(text.strip()):
            return {"text": ocr_text, "ocr_used": True, "error": None, "page_count": _pdf_page_count(file_path)}

    return {"text": text, "ocr_used": False, "error": None, "page_count": _pdf_page_count(file_path)}


def _pdf_via_pdfminer(file_path: str) -> str:
    try:
        from pdfminer.high_level import extract_text as pdfminer_extract
        return pdfminer_extract(file_path) or ""
    except Exception:
        # Fallback to PyPDF2
        try:
            import PyPDF2
            text_parts = []
            with open(file_path, "rb") as f:
                reader = PyPDF2.PdfReader(f)
                for page in reader.pages:
                    text_parts.append(page.extract_text() or "")
            return "\n".join(text_parts)
        except Exception as exc:
            logger.warning("PyPDF2 also failed: %s", exc)
            return ""


def _pdf_page_count(file_path: str) -> int:
    try:
        import PyPDF2
        with open(file_path, "rb") as f:
            reader = PyPDF2.PdfReader(f)
            return len(reader.pages)
    except Exception:
        return 0


def _try_ocr_pdf(file_path: str) -> str:
    """Attempt OCR on a PDF.  Returns empty string if pytesseract is unavailable."""
    try:
        import pytesseract
        from pdf2image import convert_from_path
        images = convert_from_path(file_path)
        parts = [pytesseract.image_to_string(img) for img in images]
        return "\n".join(parts)
    except ImportError:
        logger.info("OCR libraries not available; skipping OCR for PDF")
        return ""
    except Exception as exc:
        logger.warning("OCR failed: %s", exc)
        return ""


# ── DOCX ────────────────────────────────────────────────────────────────────

def _extract_docx(file_path: str) -> dict:
    try:
        from docx import Document
        doc = Document(file_path)
        parts = []
        for para in doc.paragraphs:
            parts.append(para.text)
        # Also grab table cells
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    parts.append(cell.text)
        return {"text": "\n".join(parts), "ocr_used": False, "error": None, "page_count": 1}
    except Exception as exc:
        return {"text": "", "ocr_used": False, "error": str(exc), "page_count": 0}


# ── TXT ─────────────────────────────────────────────────────────────────────

def _extract_txt(file_path: str) -> dict:
    for enc in ("utf-8", "latin-1", "cp1252"):
        try:
            with open(file_path, "r", encoding=enc) as f:
                return {"text": f.read(), "ocr_used": False, "error": None, "page_count": 1}
        except UnicodeDecodeError:
            continue
    return {"text": "", "ocr_used": False, "error": "Could not decode text file", "page_count": 0}


# ── Image ───────────────────────────────────────────────────────────────────

def _extract_image(file_path: str) -> dict:
    try:
        import pytesseract
        from PIL import Image
        img = Image.open(file_path)
        text = pytesseract.image_to_string(img)
        return {"text": text, "ocr_used": True, "error": None, "page_count": 1}
    except ImportError:
        return {"text": "", "ocr_used": True,
                "error": "OCR not available. Install pytesseract to process image files.",
                "page_count": 0}
    except Exception as exc:
        return {"text": "", "ocr_used": True, "error": str(exc), "page_count": 0}


# ── Text Cleaning ────────────────────────────────────────────────────────────

def clean_text(text: str) -> str:
    """Basic cleaning: normalise whitespace, remove control chars."""
    if not text:
        return ""
    # Remove non-printable chars except newlines and tabs
    text = re.sub(r"[^\x09\x0A\x0D\x20-\x7E\u00A0-\uFFFF]", " ", text)
    # Normalise line endings
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    # Collapse excessive blank lines
    text = re.sub(r"\n{3,}", "\n\n", text)
    # Collapse horizontal whitespace
    text = re.sub(r"[ \t]{2,}", " ", text)
    return text.strip()
